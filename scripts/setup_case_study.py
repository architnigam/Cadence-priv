"""Rebuild planning documents and summary PDF; does not implement the detector."""
from pathlib import Path
import shutil, hashlib, json, csv, re
from xml.sax.saxutils import escape
from pptx import Presentation
from docx import Document
from pypdf import PdfReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r'C:\Users\hp\Downloads\casestudy2\Case Study 2 Supporting Material')
def write(name, text):
    p=ROOT/name; p.parent.mkdir(parents=True,exist_ok=True); p.write_text(text.strip()+'\n',encoding='utf-8')

folders={
 'data/source':'Local copies of supplied evidence. Preserve originals.',
 'data/problem_statement':'The three supplied case study screenshots.',
 'data/extracted':'Local text, blocks, coordinates and OCR results. Generated later.',
 'data/annotations':'Human-reviewed ground truth, including repeated occurrences and image-only PII.',
 'data/synthetic':'Additional synthetic positive and negative fixtures.',
 'src/cadence_detection/ingestion':'Person 1: validation, manifest, routing and orchestration.',
 'src/cadence_detection/extraction':'Person 2: PDF OCR, DOCX and PPTX extraction.',
 'src/cadence_detection/detection':'Person 3: recognizers, classification and span resolution.',
 'src/cadence_detection/redaction':'Person 4: text masking and safe output gate.',
 'src/cadence_detection/evaluation':'Person 5: metrics and release checks.',
 'app':'Person 4: upload and results demonstration.',
 'tests/extraction':'Person 2: format and reading-order fixtures.',
 'tests/detection':'Person 3: entity and negative fixtures.',
 'tests/redaction':'Person 4: masking and leakage checks.',
 'tests/integration':'Person 5: full pipeline and failure-path checks.',
 'docs/design':'Person 1: blueprint, architecture and interface decisions.',
 'docs/team':'Five-person ownership and milestone planning.',
 'docs/review':'Source inventory and requirements traceability.',
 'presentation':'Person 5 assembles the final PPTX; every person supplies their section.',
 'output/redacted':'Generated sanitized derivatives; verify before downstream use.',
 'output/reports':'Generated evaluation reports. No measured results yet.',
 'output/pdf':'Case study summary and detailed team plan.',
 'config':'Detection taxonomy and draft pipeline settings.',
 'notebooks':'Optional local experiment or demonstration notebooks.',
}
for p,desc in folders.items():
    (ROOT/p).mkdir(parents=True,exist_ok=True)
    if not (ROOT/p/'README.md').exists(): write(p+'/README.md', '# '+p+'\n\n'+desc+'\n\nStatus: scaffold; implementation pending.')
manifest=[]
for f in sorted(SOURCE.iterdir()):
    if f.suffix.lower() not in ('.pptx','.pdf','.docx'): continue
    dest=ROOT/'data/source'/f.name
    if not dest.exists(): shutil.copy2(f,dest)
    manifest.append({'file':f.name,'source':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'bytes':f.stat().st_size})
for name in ['Screenshot 2026-10-13 135759.png','Screenshot 2026-10-13 135911.png','Screenshot 2026-10-13 135926.png']:
    f=Path(r'C:\Users\hp\Pictures\Screenshots')/name
    dest=ROOT/'data/problem_statement'/name
    if not dest.exists(): shutil.copy2(f,dest)
write('docs/review/source_manifest.json',json.dumps(manifest,indent=2))
# Retain all machine-readable source text for local reference, not as validated detector output.
parts=[]
for f in sorted((ROOT/'data/source').iterdir()):
    if f.suffix=='.pptx':
        for i,s in enumerate(Presentation(f).slides,1):
            parts.append(f'\n## {f.name} - slide {i}')
            for sh in s.shapes:
                if sh.has_text_frame: parts.append(sh.text)
                if sh.has_table:
                    parts.extend(' | '.join(c.text for c in row.cells) for row in sh.table.rows)
            if s.has_notes_slide: parts.append(s.notes_slide.notes_text_frame.text)
    elif f.suffix=='.docx':
        d=Document(f); parts.append(f'\n## {f.name} - paragraphs')
        parts.extend(p.text for p in d.paragraphs if p.text)
        for i,t in enumerate(d.tables,1):
            parts.append(f'\n### Table {i}');parts.extend(' | '.join(c.text for c in row.cells) for row in t.rows)
    elif f.suffix=='.pdf':
        parts.append(f'\n## {f.name}: 35 scanned pages; native text is effectively empty. Visual review completed; OCR implementation pending.')
write('data/extracted/source_text_reference.md','\n'.join(parts))

overview='''# Cadence Detection - Case Study 2

## Objective and current status
Build a repeatable, scalable pipeline that extracts text from PDF, PPTX and DOCX evidence, detects and classifies personally identifiable information (PII), and redacts it before any downstream LLM processing. Cadence is a global financial services firm with investment banking, asset management and consumer banking operations.

This repository currently contains a reviewed project scaffold, planning documents, source copies and a summary PDF. The detector, OCR pipeline, UI, measured evaluation results and final submission deck are not implemented yet. Person 1 through Person 5 are placeholders for the actual team members.

## What the case study requires
- A team PowerPoint (.pptx) submission and a 10-15 minute stakeholder walkthrough.
- Solution Blueprint: functional design diagram showing ingestion through PII detection output.
- System Architecture: technical components, interactions and data flow.
- Extraction Logic: extraction across formats, plus entity detection and classification logic.
- Reliable Outcomes: retain at least 80% of original text structure and preceding/succeeding context; approach 100% PII coverage while minimizing false positives. These are targets, not achieved results.
- Live Walkthrough: upload the supplied artifacts, extract their text and identify/classify PII using an IDE, notebook or simple custom UI.
- Approximately three weeks from release for the team submission. The release date was not supplied, so use relative weeks and confirm the actual deadline.

## Source review and implications
1. Organizational pack: 9 slides, named leadership, contact directory, three divisional structures and two RACI matrices. Full names, abbreviated names, corporate emails and mobile numbers occur repeatedly, including tables and org charts. Slide 6 names the DPO accountable for data classification and PII detection. Short synthetic phone formats such as +1 (555) 0114 require contextual recognition rather than only strict real-number validation. Slide 2 says five reporting divisions but lists six areas; preserve the source rather than silently correcting it.
2. Risk policy v6.0: 35 scanned pages, with effectively no native text and one page image per page. Requires OCR on every page, including smaller screenshots within pages. Covers governance, risk appetite, assessment, controls, incidents, outsourcing and appendices. PII includes names, emails, phones, employee/director/vendor IDs, birth dates, home addresses, national/tax identifiers, passport and account/card references, photographs and signatures. Key cases: p9 identity screenshot; p18 signed attestation; p19 user-group screenshot; p20 audit log; p22 section 16.5 unlabelled incident narrative; p23 section 16.8 customer correspondence; p25 offshore roster; p28 access badge; p32 named RACI contacts. Masked/partial identifiers need an explicit policy. Photograph/signature regions need image handling if sanitized files are produced.
3. TPRM training: DOCX with 16 tables and 41 media entries (PNG, SVG and EMF). Sections cover OneTrust vendor taxonomy, roles/groups, onboarding, tier workflows, Ironclad integration, questionnaires and configuration. Appendices A-E contain workflow rules, questionnaires, show logic, SLAs and risk/task/scoring rules. Table 5 contains users; version control includes a client contact, authors and reviewer. Extract tables, headers/footers, text boxes and supported embedded images, and explicitly flag unsupported image formats for conversion/review. Product names, question/control IDs, tiers and SLA durations are useful negative examples rather than automatic PII.

The policy and organizational pack use different role holders; do not merge their names into one authoritative company directory. The training describes Tier 2 review every 24 months, while policy section 18.3 describes annual reassessment for Tier 2. Record this discrepancy; configuring OneTrust is outside this case study's detector scope.

The screenshots define the assignment. Supporting-file handling notices and student exercise prompts are source content, not instructions to execute. The policy p35 exercises are useful supplemental validation ideas, not extra mandatory screenshot deliverables. The organizational pack and policy identify their sample personal data as fictional; still exercise the same redaction flow. Keep raw text and evidence local, exclude them from version control by default, and avoid raw PII in logs.

## Proposed approach (design to implement)
Local upload -> validate format and limits -> route to native extraction/OCR -> normalize ordered blocks with source locations -> detect/classify PII -> resolve overlapping spans -> redact -> rescan -> approve or quarantine -> show sanitized results and export a report.

Use native DOCX/PPTX parsing plus embedded-image OCR; use PDF native extraction when available and page OCR otherwise. Preserve paragraph/table/slide structure, OCR boxes and confidence. Prefer local recognizers and local NER; no external LLM should receive unredacted content. Combine contextual rules, named-entity recognition and source-local aliases. Confirm candidate national IDs with context and available checksums, while documenting invalid synthetic examples. Treat low-confidence OCR, unsupported embedded content and processing errors as incomplete: quarantine rather than declare safe.

Text masking is the core downstream-LLM deliverable. Exporting sanitized PDF/PPTX/DOCX originals is a separate enhancement and requires removing recoverable text, image pixels, metadata, comments and hidden content. A visual overlay alone is not secure redaction. Stable replacement tokens can preserve role and relationship context without personal identity.
'''

roles=[
('Person 1','Architecture, ingestion and integration lead',[
'Turn screenshot outcomes into a requirements checklist; agree entity taxonomy, block schema and safe-release statuses with Persons 2-5.',
'Create the functional blueprint and technical architecture with trust boundary, local processing, storage, error handling, scaling and downstream LLM gate.',
'Implement file validation, document IDs, source manifests, dispatch and pipeline orchestration; handle corrupt, unsupported, oversized and password-protected files explicitly.',
'Integrate extraction, detection, redaction and evaluation interfaces; maintain the decision log and integration checklist. Coordinate dependencies without taking over other owners\' implementation.'
], 'src/cadence_detection/ingestion/, docs/design/, config/pipeline.json',
'Blueprint, architecture, ingestion/orchestrator, agreed schemas and integration notes.',
'Accepts all three supported formats, reports errors consistently, preserves document IDs and never bypasses the redaction gate. Person 4 reviews the trust boundary; Person 5 checks integration.',
'Week 1: architecture and interface freeze. Week 2: integration pipeline. Week 3: resolve integration defects and present blueprint/architecture (about 2 minutes).'),
('Person 2','Document extraction and OCR owner',[
'Implement PDF page-level extraction/OCR for all 35 scanned policy pages; retain page number, text blocks, coordinates, reading order and OCR confidence.',
'Implement DOCX paragraph/table/header/footer/text-box extraction and PPTX slide/shape/table/group/notes extraction. Scan embedded images, and flag SVG/EMF conversion requirements.',
'Preserve row/column relationships, headings and surrounding context; add crop/upscale passes for small screenshots and handle rotation or degraded scans.',
'Create format fixtures and a page/slide extraction coverage report. Supply ordered blocks and an extraction-completeness signal to Persons 3 and 4.'
], 'src/cadence_detection/extraction/, tests/extraction/, data/extracted/',
'Three format adapters, OCR wrapper, common block output and extraction coverage report.',
'Every policy page and every PPTX slide is accounted for. Missing/unsupported content is flagged. Person 5 independently scores the agreed structure rubric; target >=80%. Person 1 reviews interfaces.',
'Week 1: native adapters and PDF OCR spike. Week 2: embedded images, tables and context. Week 3: fix hard scans and present extraction (about 2 minutes).'),
('Person 3','PII detection and classification owner',[
'Define recognizers for names/initials, emails, contextual phone formats, employee/director/vendor IDs, birth dates, addresses, national/tax IDs, passports and financial identifiers.',
'Combine local NER with rules and context; support aliases, hyphenated/international names, repeated contacts, narrative prose and partially masked identifiers.',
'Resolve overlapping spans deterministically; return type, block-local start/end, confidence, detector and location without putting raw values in routine logs.',
'Build positive and negative fixtures, including policy sections 16.5/16.8 and PPTX RACI initials. Distinguish organizations, policy/control/question IDs, dates and money from personal identifiers.'
], 'src/cadence_detection/detection/, tests/detection/, config/pii_taxonomy.json',
'Detection/classification engine, recognizer catalogue, threshold decisions and entity tests.',
'All agreed entity categories have positive and negative cases; offsets remain valid after normalization; metrics are reported by category and source type. Person 2 checks source mapping; Person 5 independently evaluates.',
'Week 1: taxonomy and baseline recognizers. Week 2: NER, aliases and narrative cases. Week 3: tune from held-out evaluation and present detection (about 2 minutes).'),
('Person 4','Redaction, privacy gate and demo application owner',[
'Implement overlap-safe replacement in descending offset order; preserve surrounding context and use stable typed placeholders within a document.',
'Build the safe-release gate: re-detect after masking, block incomplete extraction or uncertain content, and expose approved/quarantined/failed states clearly.',
'Build a local upload/results UI or notebook showing format, extraction status, entity categories, source locations, sanitized text and downloadable safe output.',
'Ensure raw text, temporary images and identity mappings stay local; define cleanup and log rules. Document text-output scope and any unimplemented original-file redaction limitations.'
], 'src/cadence_detection/redaction/, tests/redaction/, app/, output/redacted/',
'Text redactor, release gate, local demo interface, safe exports and demonstration runbook.',
'No raw spans survive the known-fixture checks; gate blocks OCR/adapter failures; repeated tokens are consistent and non-PII context remains readable. Person 3 reviews masking; Person 5 runs leakage/failure tests.',
'Week 1: masking and mock UI using schema fixtures. Week 2: connect real pipeline and gate. Week 3: polish demo/rehearse and demonstrate all three files (about 4 minutes).'),
('Person 5','Independent evaluation, QA and submission owner',[
'Create and adjudicate the ground truth with all four teammates contributing annotations. Label every occurrence, source location, entity type and image-only status; reserve a held-out subset before tuning.',
'Measure precision, recall, F1, false positives/negatives, per-format and per-category performance, OCR/image-only recall and remaining PII after redaction.',
'Define the 80% structure rubric before implementation; score headings, reading order, paragraph boundaries, table relationships and neighboring context. Keep structure score separate from text/OCR accuracy.',
'Run full pipeline, corrupt/unsupported/empty-file and quarantine tests. Assemble final PPTX from each owner\'s content, evidence and diagrams; coordinate a 10-15 minute rehearsal and submission checklist.'
], 'src/cadence_detection/evaluation/, tests/integration/, data/annotations/, output/reports/, presentation/',
'Annotation guide and gold set, evaluation script/report, defect tracker, final stakeholder PPTX and demo fallback material.',
'Every mandatory outcome maps to evidence. Reports give denominators, unresolved cases and measured results rather than claiming universal 100% detection. Person 1 reviews deck completeness; Person 2 checks structure scoring.',
'Week 1: annotation guide, rubric and test plan. Week 2: baseline report and deck outline. Week 3: final regression, metrics, deck assembly and reliability/results presentation (about 2 minutes).')]

team='## Detailed five-person work allocation\n\nEach person owns an implementation component, its focused tests, documentation and a presentation contribution. Person 5 coordinates QA and assembly; everyone supplies reviewed content and annotations. Replace the placeholders with names when the team is known.\n\n'
team+='| Person | Main ownership | Primary handoff |\n|---|---|---|\n'
for who,title,tasks,paths,deliver,done,week in roles: team+=f'| {who} | {title} | {deliver} |\n'
for who,title,tasks,paths,deliver,done,week in roles:
    team+=f'\n### {who}: {title}\n\n**Responsibilities**\n\n'+''.join(f'- {t}\n' for t in tasks)
    team+=f'\n**Owned paths:** {paths}\n\n**Deliverables:** {deliver}\n\n**Definition of done and review:** {done}\n\n**Weekly plan and presentation:** {week}\n'

closing='''## Handoffs and three-week coordination
- Days 1-2: Person 1 agrees schemas; Person 5 agrees labels, split and metric definitions. Every person reviews the source inventory.
- Week 1 checkpoint: Person 2 supplies extraction fixtures; Person 3 supplies detections on fixtures; Person 4 supplies a masking/UI skeleton; Person 5 supplies an annotation baseline. Architecture and interfaces are agreed together.
- Week 2 checkpoint: Person 1 runs all three files through the integrated pipeline. Person 5 publishes baseline results and defects; Persons 2-4 fix their components. All owners contribute slide drafts.
- Week 3 checkpoint: freeze features early, rerun evaluation, resolve blockers, assemble PPTX and rehearse. Person 4 prepares the local live demo; Person 5 prepares screenshots/recorded fallback. Confirm the actual submission date separately.
- Dependency chain: ingestion -> extraction -> detection -> redaction/gate -> evaluation/demo. Use schema fixtures to let Persons 3-5 work before OCR is complete. Person 5's gold set must not be replaced by Person 3's own predictions.

## Evaluation and completion criteria
- Entity precision = TP/(TP+FP), recall = TP/(TP+FN), F1 = 2PR/(P+R). Use exact occurrence/type matching as primary and separately report partial matches and OCR errors. Count repeated mentions separately.
- Proposed structure score: 20% headings, 20% reading order, 20% paragraph/block boundaries, 20% table row/column relationships, 20% neighboring context. Human reviewers use the same annotated units; mark non-applicable dimensions and renormalize. Target >=80% per document, reported alongside OCR/text accuracy. This rubric is a team proposal because the screenshots do not define a formula.
- Report native-text versus image-only recall on additional fixtures; all supplied policy pages are scans, so native PDF comparison needs another fixture or controlled transcription baseline.
- Ground truth includes p22 narrative, p23 correspondence, p18 attestation, p19 user screenshot, p20 audit log, p25 roster, p28 badge and PPTX abbreviated names. Use synthetic negative cases for PAN/PESEL/TIN/card patterns, policy IDs and generic dates.
- Re-scan sanitized text; independently check known gold spans and a sample manually. A zero residual count on this dataset is evidence for that test only, not a guarantee for future documents.
- Preserve source locations for each finding and for each failure. No silent dropped pages/images. Quarantine incomplete work instead of forwarding it to an LLM.

## Repository layout
```text
Cadence_Detection/
  README.md
  config/                         draft pipeline and entity taxonomy
  data/source/                    supplied evidence copies (local)
  data/problem_statement/         supplied screenshots (local)
  data/extracted/                 reference text; future OCR blocks (local)
  data/annotations/               annotation schema; future gold set (local)
  data/synthetic/                 future additional test inputs
  src/cadence_detection/
    ingestion/ extraction/ detection/ redaction/ evaluation/
  app/                            future local demo
  tests/                          extraction, detection, redaction, integration
  docs/design/                    proposed diagrams and interface contract
  docs/review/                    requirements and source inventory
  docs/team/                      detailed work allocation
  presentation/                   deck outline; final PPTX to be created
  output/pdf/                     summary PDF
  output/reports/                  future evaluation results
  output/redacted/                 future safe exports
  scripts/                        reproducible planning/PDF generation
  notebooks/                      optional experiments
```

## Working rules and first actions
1. Add actual names next to Person 1-5; keep ownership and review responsibilities explicit.
2. Read docs/review/source_review.md and docs/design/interfaces.md before coding.
3. Person 1 finalizes interfaces with everyone; Persons 2-4 start with their fixtures; Person 5 establishes independent labels.
4. Choose and pin runtime/OCR/model dependencies after the local prototype is agreed. requirements-planned.txt is a shortlist, not a working installation lockfile.
5. Never treat the reference text dump as OCR output or validated ground truth. Run the detector only after it is implemented; there is no runnable detection command yet.
6. Each owner updates progress, evidence, open defects and slide material. Do not commit raw source files, extracted PII, identity mappings or local caches by default.

## Deliverable status
- Completed: source review, source manifest/copies, scaffold, draft contracts/design, work allocation, summary PDF and deck outline.
- Pending: extraction/OCR, PII recognition, redaction, demo, annotations, measured results, final PPTX and stakeholder rehearsal.
'''
readme=overview+'\n'+team+'\n'+closing
write('README.md',readme)
write('docs/team/work_allocation.md','# Five-person team plan\n\n'+team+'\n'+closing.split('## Repository layout')[0])
write('docs/review/source_review.md','# Source review\n\n'+overview.split('## Source review and implications')[1].split('## Proposed approach')[0])
write('docs/review/requirements.md','''# Requirement traceability

| Screenshot outcome | Owner | Evidence to produce |
|---|---|---|
| Solution Blueprint | Person 1 | Functional diagram and narrated data flow |
| System Architecture | Person 1, all component owners | Technical diagram and trust boundary |
| Extraction Logic | Persons 2 and 3 | Format adapters, OCR and recognizer explanation |
| Reliable Outcomes | Person 5, Persons 2-4 | Structure rubric, gold labels, metrics and leakage checks |
| Live Walkthrough | Person 4, Persons 1 and 5 | Upload all formats, sanitized results and failure handling |
| PPTX; 10-15 minutes | Person 5; all contribute | Final deck and timed rehearsal |

Targets: >=80% structure/context preservation; approach 100% PII coverage with minimal false positives. Actual deadline: confirm release date + approximately three weeks.
''')
write('docs/design/interfaces.md','''# Draft component contract

Agree before implementation. All statuses below are proposals.

- Document: document_id, filename, detected_format, sha256, processing_status, errors.
- Block: document_id, block_id, page_or_slide (1-based), source_kind, reading_order, text, table_row/table_column (optional), bbox (optional), OCR_confidence (optional), extraction_complete.
- Bounding boxes: specify coordinate system and units per adapter; retain conversion to original image/slide/page. Never mix normalized and pixel coordinates silently.
- Finding: document_id, block_id, entity_type, start, end, confidence, detector, source_location. Python Unicode code-point offsets; start inclusive/end exclusive into the exact normalized block text. Raw values stay in memory/local restricted evidence only.
- Redaction: resolve overlaps, replace from right to left, generate stable document-local typed tokens and preserve block structure. Mapping storage requires explicit local controls.
- Result: approved, quarantined or failed; sanitized_blocks, category_counts, extraction_coverage, remaining_candidate_count, warnings. Approval requires completed extraction and a passed verification gate; gate criteria to be agreed.
- Implement adapters independently of UI. Person 1 owns contract changes with all consumers reviewing; fixtures should cover errors and success.
''')
write('docs/design/blueprint.md','''# Proposed functional blueprint

```mermaid
flowchart LR
 U[Local upload] --> V[Validate and route]
 V --> E[Native extraction and local OCR]
 E --> B[Ordered blocks and source locations]
 B --> D[Local rules and NER]
 D --> R[Classify and redact]
 R --> G[Rescan and completeness gate]
 G -->|pass| S[Sanitized text and report]
 G -->|incomplete or uncertain| Q[Quarantine and review]
 S --> L[Downstream LLM boundary]
```

Raw evidence remains inside the local processing boundary. Future scale: job queue, bounded OCR workers and document-level isolation. Final technical architecture must specify actual selected libraries, storage, deployment and failure recovery; these are not implemented.
''')
write('presentation/deck_outline.md','''# Final submission outline (PPTX pending)

Target 12 minutes plus up to 3 minutes of discussion.

1. Problem, scope and functional blueprint - Person 1.
2. Technical architecture and privacy boundary - Person 1. Total: 2 min.
3. Format extraction/OCR and structure preservation - Person 2. 2 min.
4. Entity taxonomy, detection logic and hard cases - Person 3. 2 min.
5. Measured precision/recall, structure score and limitations - Person 5. 2 min.
6. Live upload/extract/classify/redact for all three files - Person 4. 4 min.
7. Summary, open risks and questions - shared within remaining time.

Every owner supplies reviewed slide content and evidence. Person 5 assembles consistent slides. Include both functional and detailed technical diagrams; do not substitute this outline or the planning PDF for the required final PowerPoint.
''')
write('config/pii_taxonomy.json',json.dumps({'status':'proposed','types':['PERSON','EMAIL','PHONE','EMPLOYEE_ID','DIRECTOR_ID','VENDOR_PERSON_ID','DATE_OF_BIRTH','HOME_ADDRESS','NATIONAL_ID','TAX_ID','PASSPORT','BANK_ACCOUNT','PAYMENT_CARD','PARTIAL_IDENTIFIER','PHOTOGRAPH_REGION','SIGNATURE_REGION'],'notes':'Dates and addresses require personal context. Visual regions need separate image handling. Confirm scope before implementation.'},indent=2))
write('config/pipeline.json',json.dumps({'status':'draft_not_implemented','formats':['pdf','pptx','docx'],'processing':'local','ocr':'required_for_scans_and_embedded_images','failure_policy':'quarantine','downstream_input':'verified_sanitized_text_only'},indent=2))
write('data/annotations/annotation_schema.json',json.dumps({'document_id':'','location':{'page_or_slide':None,'block_id':'','bbox':None},'entity_type':'','start':None,'end':None,'source_kind':'body|table|narrative|image|header|footer|notes','image_only':False,'reviewer':'','adjudication_status':'pending','notes':''},indent=2))
write('requirements-planned.txt','''# Candidate dependencies; choose versions and OCR/model assets during implementation.
python-docx
python-pptx
pypdf
Pillow
reportlab
# Consider: a local OCR engine plus layout support, local NER/PII recognizers,
# a local UI framework, and pytest. Not installed or pinned by this scaffold.
''')
write('.gitignore','''__pycache__/
*.py[cod]
.venv/
.env
tmp/
data/source/*
!data/source/README.md
data/problem_statement/*
!data/problem_statement/README.md
data/extracted/*
!data/extracted/README.md
data/annotations/*
!data/annotations/README.md
!data/annotations/annotation_schema.json
output/redacted/*
!output/redacted/README.md
output/reports/*
!output/reports/README.md
''')

# Use the README as the content source so the PDF and allocation remain consistent.
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyCase',fontName='Helvetica',fontSize=9.3,leading=13,spaceAfter=7,textColor=colors.HexColor('#243449')))
styles['Heading1'].textColor=colors.HexColor('#143052');styles['Heading2'].textColor=colors.HexColor('#143052')
story=[]
story.append(Paragraph('Cadence | Case Study 2',styles['Title']))
story.append(Paragraph('Text analysis and PII detection',styles['Heading2']))
story.append(Paragraph('Case summary, implementation scaffold and detailed five-person delivery plan',styles['BodyCase']))
story.append(Paragraph('Prepared 13 October 2026',styles['BodyCase']))
story.append(Spacer(1,16))
in_code=False
for line in readme.splitlines()[2:]:
    if line.startswith('```'): in_code=not in_code;continue
    if in_code or not line.strip() or line.startswith('|'): continue
    if line == '## Repository layout': continue
    if line.startswith('## '):
        if line in ['## Detailed five-person work allocation','## Handoffs and three-week coordination','## Working rules and first actions']: story.append(PageBreak())
        story.append(Paragraph(escape(line[3:]),styles['Heading1']))
    elif line.startswith('### '):
        if not line.startswith('### Person 1:'): story.append(PageBreak())
        story.append(Paragraph(escape(line[4:]),styles['Heading2']))
    else:
        t=escape(line)
        t=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',t)
        if t.startswith('- '): t='- '+t[2:]
        story.append(Paragraph(t,styles['BodyCase']))
def footer(canvas,doc):
    canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#68778a'))
    canvas.drawString(44,27,'Cadence Detection | Case summary and team plan')
    canvas.drawRightString(551,27,str(doc.page))
out=ROOT/'output/pdf/Cadence_Case_Study_Summary_and_Team_Plan.pdf'
SimpleDocTemplate(str(out),pagesize=(595,842),rightMargin=44,leftMargin=44,topMargin=42,bottomMargin=44).build(story,onFirstPage=footer,onLaterPages=footer)
print('Created',out,'pages',len(PdfReader(out).pages))
print('Created scaffold and synchronized README/team plan/PDF.')
